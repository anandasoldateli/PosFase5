"""Gera o documento Word da síntese integrada (Q1–Q11).

Todos os números vêm de ``outputs/resultados_analises.json`` e
``outputs/resultados_modelo.json``; as figuras vêm de ``outputs/figuras``.
Rankings e comparações citados no texto são calculados a partir dos dados;
o texto interpretativo usa linguagem de associação, nunca de causalidade.

Uso: python documento.py  →  outputs/Datathon_Passos_Magicos_Sintese_Q1_Q11.docx
"""

import json
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from config import PASTA_FIGURAS, PASTA_OUTPUTS, RESULTADOS_ANALISES, RESULTADOS_MODELO

SAIDA = PASTA_OUTPUTS / "Datathon_Passos_Magicos_Sintese_Q1_Q11.docx"

AZUL = RGBColor(0x1F, 0x4E, 0x79)
CINZA = RGBColor(0x64, 0x64, 0x64)
FUNDO_CAIXA = "EAF2F8"
FUNDO_CABECALHO = "D9EAF7"
FONTE = "Aptos"


# ---------------------------------------------------------------------------
# Formatação
# ---------------------------------------------------------------------------

def br(v, casas=1):
    return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(v, casas=1):
    return f"{br(v, casas)}%"


def fmt_p(p):
    return "p < 0,001" if p < 0.001 else f"p = {br(p, 3)}"


def ic(par, casas=3):
    return f"{br(par[0], casas)}–{br(par[1], casas)}"


def fase(f):
    f = int(f)
    return "ALFA" if f == 0 else f"Fase {f}"


def sombrear(celula, cor):
    tc_pr = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    tc_pr.append(sombra)


def configurar_estilos(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    normal.paragraph_format.space_after = Pt(4)
    for nome, tamanho in [("Title", 24), ("Subtitle", 13), ("Heading 1", 15), ("Heading 2", 12.5)]:
        estilo = doc.styles[nome]
        estilo.font.name = FONTE
        estilo.font.size = Pt(tamanho)
        estilo.font.color.rgb = AZUL
        estilo.font.bold = nome != "Subtitle"
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    for secao in doc.sections:
        secao.left_margin = secao.right_margin = Cm(2.2)
        secao.top_margin = secao.bottom_margin = Cm(2.0)


class Documento:
    def __init__(self):
        self.doc = Document()
        configurar_estilos(self.doc)
        # O template padrão do python-docx omite w:percent, exigido pelo esquema OOXML.
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None:
            zoom.set(qn("w:percent"), "100")
        self.n_figura = 0
        self.n_tabela = 0

    def titulo(self, texto, estilo="Heading 1"):
        self.doc.add_paragraph(texto, style=estilo)

    def par(self, texto, negrito_prefixo=None, italico=False):
        p = self.doc.add_paragraph()
        if negrito_prefixo:
            p.add_run(negrito_prefixo).bold = True
        p.add_run(texto).italic = italico
        return p

    def pergunta(self, texto):
        p = self.doc.add_paragraph()
        r = p.add_run("Pergunta: ")
        r.bold = True
        r.font.color.rgb = AZUL
        p.add_run(texto).italic = True

    def resposta(self, texto):
        p = self.doc.add_paragraph()
        r = p.add_run("Resposta: ")
        r.bold = True
        r.font.color.rgb = AZUL
        p.add_run(texto).bold = True

    def itens(self, lista):
        for item in lista:
            p = self.doc.add_paragraph(style="List Bullet")
            if isinstance(item, tuple):
                p.add_run(item[0]).bold = True
                p.add_run(item[1])
            else:
                p.add_run(item)

    def caixa(self, rotulo, texto):
        tabela = self.doc.add_table(rows=1, cols=1)
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        celula = tabela.rows[0].cells[0]
        sombrear(celula, FUNDO_CAIXA)
        p = celula.paragraphs[0]
        r = p.add_run(f"{rotulo}. ")
        r.bold = True
        r.font.color.rgb = AZUL
        p.add_run(texto)
        self.doc.add_paragraph()

    def figura(self, arquivo, legenda, largura=15.5):
        self.n_figura += 1
        self.doc.add_picture(str(PASTA_FIGURAS / f"{arquivo}.png"), width=Cm(largura))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Figura {self.n_figura} — {legenda}")
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = CINZA

    def tabela(self, cabecalho, linhas, larguras=None, legenda=None):
        tabela = self.doc.add_table(rows=1, cols=len(cabecalho))
        tabela.style = "Table Grid"
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        for celula, texto in zip(tabela.rows[0].cells, cabecalho):
            sombrear(celula, FUNDO_CABECALHO)
            r = celula.paragraphs[0].add_run(str(texto))
            r.bold = True
            r.font.size = Pt(9.5)
        for linha in linhas:
            celulas = tabela.add_row().cells
            for celula, texto in zip(celulas, linha):
                r = celula.paragraphs[0].add_run(str(texto))
                r.font.size = Pt(9.5)
        if larguras:
            for linha in tabela.rows:
                for celula, largura in zip(linha.cells, larguras):
                    celula.width = Cm(largura)
        if legenda:
            self.n_tabela += 1
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"Tabela {self.n_tabela} — {legenda}")
            r.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = CINZA
        else:
            self.doc.add_paragraph()

    def quebra(self):
        self.doc.add_page_break()

    def salvar(self, caminho):
        self.doc.save(caminho)


# ---------------------------------------------------------------------------
# Conteúdo
# ---------------------------------------------------------------------------

def construir(a, m):
    d = a["diagnostico"]
    q1, q2, q3, q4, q5, q6, q7, q8, q10, q11 = (a[k] for k in ["Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7", "Q8", "Q10", "Q11"])
    anos = ["2022", "2023", "2024"]
    t = m["teste_modelo_escolhido"]
    ma, mm = m["metricas_alto"], m["metricas_medio"]
    top10, top20 = m["captura_top"][0], m["captura_top"][1]
    nr = q11["nao_reaparecimento"]
    nrp = q10["nao_reaparecimento_por_pedra"]
    tp = q11["taxa_promocao_por_fase"]
    npromo = q11["n_promocao_por_fase"]
    g24 = q8["potencial_aritmetico"]["2024"]["ganho_inde_se_10"]
    maior_potencial = max(g24, key=g24.get)
    perfil = q8["perfis"]["2024"]["por_numero_de_indicadores_altos"]
    coorte = q10["coorte_por_pedra_inicial"]
    cf, cr = q1["coorte_pct_defasagem_negativa"], q1["coorte_pct_defasagem_negativa_regra"]
    s_regra = m["sensibilidade_regra_2024"]
    sens = {v["variante"]: v for v in m["sensibilidade_atributos"]}
    anterior = next(v for k, v in sens.items() if k.startswith("Alvo anterior"))
    faixa_teste = [v["ROC-AUC (teste)"] for k, v in sens.items() if not k.startswith("Alvo anterior")]
    rob = {r["análise"]: r for r in m["robustez"]}
    dec = m["decomposicao_evento"]
    pxe = m["promocao_x_evento"]
    comp_alto = m["composicao_faixa_alto"]
    n_alto = sum(comp_alto.values())
    calib = {r["fase"]: r for r in m["calibracao_por_fase"]}
    pior_calib = sorted(calib.values(), key=lambda r: r["diferença (p.p.)"])[:2]
    final = m["modelo_final"]
    fases_q2 = q2["fases_exibidas"]
    rank = q2["ranking_fases_exibidas"]
    menor_sempre = {rank[a]["menor"] for a in anos}
    maior_sempre = {rank[a]["maior"] for a in anos}

    doc = Documento()

    # Capa
    p = doc.doc.add_paragraph("DATATHON — PASSOS MÁGICOS", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.doc.add_paragraph("Síntese integrada das análises — Diagnóstico e Questões 1 a 11", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for texto in [f"Base longitudinal PEDE 2022–2024 · Versão revisada · {date.today():%d/%m/%Y}",
                  "Ananda Soares Soldateli · Grecco Teixeira de Morais"]:
        p = doc.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(texto).font.color.rgb = CINZA

    doc.caixa("OBJETIVO", "Responder às 11 perguntas do Datathon com base nos dados da PEDE 2022–2024, apresentar o modelo "
              "preditivo de risco e propor oportunidades de melhoria à Associação. Todos os números são gerados pelo "
              "pipeline do repositório (src/ e notebooks/), sem valores digitados à mão. Os resultados são observacionais: "
              "mostram associações, não efeitos causais do programa.")

    # Resumo executivo
    doc.titulo("Resumo executivo")
    doc.tabela(["Tema", "Mensagem gerencial"], [
        ["Q1 — Defasagem (IAN)", f"Na fonte, os alunos em fase passam de {pct(q1['percentual']['2022']['Em fase'])} (2022) para "
         f"{pct(q1['percentual']['2024']['Em fase'])} (2024). A magnitude em 2024 depende da regra de fase ideal, que mudou nesse ano."],
        ["Q2 — Desempenho (IDA)", f"Sem tendência consistente: {br(q2['media']['2022'], 2)} → {br(q2['media']['2023'], 2)} → {br(q2['media']['2024'], 2)}. "
         f"A {fase(min(menor_sempre))} tem a menor média nos três anos."],
        ["Q3 — Engajamento (IEG)", "Associação positiva e moderada com IDA e IPV nos três anos."],
        ["Q4 — Autoavaliação (IAA)", "Associação fraca com IDA e IEG. O IAA mede percepção e bem-estar amplos, não a nota; a comparação é exploratória."],
        ["Q5 — Psicossocial (IPS)", "Não foi detectada associação entre IPS e queda de IDA ou IEG no ano seguinte."],
        ["Q6 — Psicopedagógico (IPP)", "IPP ligeiramente menor entre defasados: mesma direção do IAN, com diferença pequena."],
        ["Q7 — Ponto de virada (IPV)", "Mais associado a IDA e IEG e, a partir de 2023, ao IPP."],
        ["Q8 — INDE", f"A fórmula oficial foi reproduzida. O maior espaço aritmético está no {maior_potencial} (+{br(g24[maior_potencial], 2)} se chegasse a 10)."],
        ["Q9 — Modelo de risco", f"Prevê a redução de D no ano seguinte com ROC-AUC {br(t['roc_auc'], 3)} (IC 95% {ic(t['roc_auc_ic95'])}). "
         f"Os 10% de maior risco concentram {pct(top10['% do total de eventos'], 0)} dos eventos. Uso indicado: priorização."],
        ["Q10 — Pedras", "A distribuição melhora, mas nos mesmos alunos o INDE fica estável, exceto no Quartzo. Critérios de Pedra variam entre anos."],
        ["Q11 — Insights", f"Não reaparecimento na base concentrado nas Pedras mais baixas ({pct(nrp['2022->2023']['pct']['Quartzo'], 0)} dos Quartzo "
         f"vs {pct(nrp['2022->2023']['pct']['Topázio'], 0)} dos Topázio em 2022→23) e promoção baixa no ALFA e, em 2023→24, nas Fases 5 e 7."],
    ], larguras=[4.2, 12.4])
    doc.caixa("LEITURA GERAL", "Os dados mostram menos alunos defasados ao longo dos anos, mas o desempenho acadêmico oscila, "
              "a regra de fase ideal mudou em 2024 e a composição da base muda de um ano para o outro. As oportunidades mais "
              "claras são priorizar alunos com maior risco de perder adequação de nível, reforçar o IDA nas fases de menor "
              "desempenho e investigar por que alunos de menor INDE deixam de aparecer na base.")

    doc.titulo("Nota sobre esta versão", "Heading 2")
    doc.par("Esta versão substitui a síntese anterior (Q1–Q10) e incorpora uma auditoria independente. As principais mudanças:")
    doc.itens([
        ("Dados de 2023 recuperados: ", "INDE e Pedra estão nas colunas INDE 2023/Pedra 2023; a versão anterior lia colunas vazias."),
        ("Problema preditivo reformulado: ", "o alvo passou a ser a redução de D no ano seguinte, e as variáveis categóricas foram harmonizadas "
         f"entre anos. O alvo anterior, com os atributos corrigidos e no mesmo protocolo, atinge ROC-AUC {br(anterior['ROC-AUC (teste)'], 3)}; "
         "por isso as métricas antigas e novas não são comparáveis como ganho de modelagem."),
        ("Comparabilidade explicitada: ", "regra de fase ideal de 2024, zeros de IEG e IAA, critérios de Pedra e ausência na base "
         "passaram a ter tratamento documentado e análises de sensibilidade."),
        ("Linguagem: ", "conclusões reescritas como associação, potencial aritmético ou não detecção, conforme o que cada análise sustenta."),
    ])
    doc.quebra()

    # Diagnóstico
    doc.titulo("Diagnóstico da base de dados")
    doc.par("Foram analisadas as três abas da planilha PEDE 2024 (PEDE2022, PEDE2023 e PEDE2024), combinadas em uma base aluno × ano. "
            "As análises comparam fotografias anuais; não há dados de trajetória dentro de cada ano.")
    doc.tabela(["Ano", "Alunos", "Com INDE e Pedra", "Cobertura IDA", "IEG efetivo*", "Cobertura IPP"], [
        [a, d["registros_por_ano"][a], sum(q10["distribuicao_n"][a].values()), pct(d["cobertura_pct"][a]["IDA"]),
         pct(d["cobertura_pct"][a]["IEG_efetivo"]), "—" if a == "2022" else pct(d["cobertura_pct"][a]["IPP"])] for a in anos
    ], larguras=[2, 2.2, 3.2, 2.8, 2.8, 2.8],
        legenda=f"Registros por ano ({d['alunos_distintos']} alunos distintos; {d['alunos_3_anos']} presentes nos três anos). "
                "*IEG excluindo zeros sem nenhuma outra avaliação")
    doc.titulo("Ocorrências de qualidade e tratamento", "Heading 2")
    dif24 = d["fase_ideal_2024_diferenca"]
    doc.itens([
        ("Códigos diferentes entre anos: ", "fase (7, FASE 7, 7A), gênero (Menina/Feminino) e escola foram harmonizados."),
        ("Fase ideal em 2024: ", f"em 2022 e 2023 a fase ideal segue a mesma correspondência com a idade; em 2024, "
         f"{d['fase_ideal_diverge_regra']['2024']} registros divergem dela ({dif24.get('-1', 0)} com fase ideal uma abaixo e "
         f"{dif24.get('1', 0)} uma acima). Nenhuma data de referência única reproduz a regra de 2024. A fonte é mantida e a "
         "diferença é tratada como sensibilidade (Q1 e Q9)."),
        ("IEG 2024: ", f"{d['ieg_zero_2024']} zeros, dos quais {d['ieg_sem_avaliacao_2024']} em registros sem nenhuma outra avaliação "
         f"(Fases 8 e 9). A média muda de {br(d['ieg_medio_2024']['todos'], 2)} para {br(d['ieg_medio_2024']['sem_nao_avaliados'], 2)} "
         "sem esses registros, que são tratados como não avaliados nas análises."),
        ("IAA = 0: ", ", ".join(f"{k}: {v}" for k, v in d["iaa_zero"].items()) + " registros. Como ocorrem com as demais avaliações "
         "presentes, a hipótese é de não resposta (a confirmar com a Associação); são excluídos quando o IAA é analisado e mantidos no INDE oficial."),
        ("IPS 2023: ", f"distribuição concentrada ({d['ips_moda_2023'].get('2.52')} alunos com 2,52; média {br(q5['ips_medio']['2023'], 2)} "
         "contra ~6,9 nos outros anos). A causa (mudança de instrumento?) não está confirmada; comparações de nível entre anos são evitadas."),
        ("Outros: ", f"sem IPP em 2022; {d['inde_incluir_2024']} alunos com INDE 2024 = “INCLUIR”; "
         f"{d['defasagem_inconsistente_fonte']} registros com Defasagem ≠ Fase − Fase ideal (mantidos)."),
    ])
    doc.caixa("CRITÉRIO", "A planilha original não é alterada. Os valores da fonte são preservados e os tratamentos analíticos ficam "
              "em colunas separadas. Ausência não é tratada como zero, e as tabelas informam o número de alunos usados.")
    doc.quebra()

    # Q1
    doc.titulo("Q1 — Adequação do nível (IAN)")
    doc.pergunta("Qual é o perfil geral de defasagem dos alunos (IAN) e como ele evolui ao longo do ano?")
    c22, c24 = q1["contagem"]["2022"], q1["contagem"]["2024"]
    doc.resposta(f"A parcela em fase aumenta entre os anos: {pct(q1['percentual']['2022']['Em fase'])} em 2022, "
                 f"{pct(q1['percentual']['2023']['Em fase'])} em 2023 e {pct(q1['percentual']['2024']['Em fase'])} em 2024 (fonte). "
                 "A queda de 2023 para 2024 é menor quando se aplica a regra de fase ideal dos anos anteriores.")
    doc.figura("q1_classes_ian", "Classificação do IAN por ano na fonte (em fase: IAN 10; moderada: IAN 5; severa: IAN 2,5).")
    doc.itens([
        f"2022: {c22['Moderada']} alunos com defasagem moderada e {c22['Severa']} com defasagem severa; 2024: {c24['Moderada']} e {c24['Severa']}.",
        "IAN médio: " + " → ".join(br(q1["ian_medio"][x], 2) for x in anos) + ".",
        f"Nos mesmos {q1['coorte_n']} alunos presentes nos três anos, a parcela com D < 0 na fonte é "
        f"{pct(cf['2022'])} → {pct(cf['2023'])} → {pct(cf['2024'])}.",
    ])
    doc.figura("q1_coorte_sensibilidade", "Mesmos alunos: parcela defasada pela fonte e pela regra idade → fase de 2022/23 aplicada a 2024.")
    doc.caixa("SENSIBILIDADE", f"Aplicando em 2024 a mesma correspondência idade → fase ideal de 2022 e 2023, a parcela defasada da coorte "
              f"em 2024 seria {pct(cr['2024'])}, e não {pct(cf['2024'])}. A redução entre 2022 e 2023 não depende dessa escolha. "
              "Como a regra efetiva de 2024 não foi documentada, a magnitude da melhora em 2024 deve ser apresentada com essa ressalva.")
    doc.caixa("CUIDADO", "A pergunta fala em evolução “ao longo do ano”, mas a base traz uma fotografia por ano; a análise é entre anos. "
              f"O IAN é derivado da defasagem (correlação ≈ {br(q1['corr_ian_defasagem']['2024'], 2)}), então os dois medem a mesma coisa em escalas diferentes.")

    # Q2
    doc.titulo("Q2 — Desempenho acadêmico (IDA)")
    doc.pergunta("O desempenho acadêmico médio (IDA) está melhorando, estagnado ou caindo ao longo das fases e anos?")
    v1, v2 = q2["variacao_individual"]["2022->2023"], q2["variacao_individual"]["2023->2024"]
    excl = q2["fases_excluidas_n_minimo"]
    doc.resposta("Oscila, sem tendência de melhora: sobe em 2023 e cai em 2024. "
                 f"Entre as fases exibidas, a {fase(min(menor_sempre))} tem a menor média em todos os anos.")
    doc.figura("q2_ida_fase_ano", f"IDA médio por fase e ano (fases com pelo menos 15 alunos com IDA em todos os anos: {', '.join(fase(f) for f in fases_q2)}).")
    pf, nf = q2["por_fase"], q2["n_por_fase"]
    itens = [
        "IDA médio: " + " → ".join(br(q2["media"][x], 2) for x in anos) + "; IDA ≤ 5: " + " → ".join(pct(q2["pct_ate_5"][x]) for x in anos) + ".",
        f"Mesmos alunos 2022→2023 (n = {v1['n']}): variação média {br(v1['delta_medio'], 2)}, mediana {br(v1['delta_mediano'], 2)}; "
        f"o teste de Wilcoxon de postos sinalizados não indica deslocamento das variações ({fmt_p(v1['p_wilcoxon'])}).",
        f"Mesmos alunos 2023→2024 (n = {v2['n']}): variação média {br(v2['delta_medio'], 2)}, mediana {br(v2['delta_mediano'], 2)}; "
        f"{pct(v2['pct_piorou'])} pioraram; Wilcoxon {fmt_p(v2['p_wilcoxon'])}.",
    ]
    if len(maior_sempre) == 1:
        f_maior = next(iter(maior_sempre))
        itens.append(f"Entre as fases exibidas, o maior IDA é sempre o da {fase(f_maior)} e o menor, o da {fase(min(menor_sempre))}: "
                     + "; ".join(f"{x}: {br(pf[str(f_maior)][x], 2)} vs {br(pf[str(min(menor_sempre))][x], 2)}" for x in anos) + ".")
    for f in excl:
        itens.append(f"A {fase(f)} ficou fora do gráfico por ter poucos alunos com IDA ("
                     + ", ".join(f"{x}: n = {nf[str(f)][x]}" for x in anos) + "); em 2023 sua média foi "
                     f"{br(pf[str(f)]['2023'], 2)}, com n = {nf[str(f)]['2023']}.")
    doc.itens(itens)
    doc.caixa("LEITURA", f"O desempenho acadêmico não acompanha a redução da defasagem. A {fase(min(menor_sempre))} (7º e 8º ano) concentra "
              "as menores médias e é um candidato natural a reforço de Matemática e Português.")

    # Q3
    doc.titulo("Q3 — Engajamento (IEG)")
    doc.pergunta("O grau de engajamento dos alunos (IEG) tem relação direta com seus indicadores de desempenho (IDA) e do ponto de virada (IPV)?")
    doc.resposta("Há associação positiva e moderada, estável nos três anos, tanto com o IDA quanto com o IPV.")
    doc.figura("q3_ieg_quartis", "IDA e IPV médios por quartil de IEG (quartis por valor, entre alunos com IEG, IDA e IPV).")
    doc.tabela(["Ano", "IEG médio*", "r (IEG, IDA)", "n", "r (IEG, IPV)", "n"],
               [[x, br(q3["ieg_medio"][x], 2), br(q3["corr"][x]["IEG_IDA"]["r"], 3), q3["corr"][x]["IEG_IDA"]["n"],
                 br(q3["corr"][x]["IEG_IPV"]["r"], 3), q3["corr"][x]["IEG_IPV"]["n"]] for x in anos],
               larguras=[2, 2.6, 2.6, 2, 2.6, 2],
               legenda="Correlações de Pearson (todas com p < 0,001). *Sem os zeros de registros não avaliados")
    qq = q3["por_quartil_ieg"]["2024"]
    lim = q3["limites_quartis"]["2024"]
    doc.itens([f"Em 2024, o IDA médio vai de {br(qq['Q1']['IDA'], 2)} no quartil de menor IEG (até {br(lim[1], 2)}; n = {int(qq['Q1']['n'])}) "
               f"a {br(qq['Q4']['IDA'], 2)} no de maior (acima de {br(lim[3], 2)}; n = {int(qq['Q4']['n'])}); o IPV, de {br(qq['Q1']['IPV'], 2)} a {br(qq['Q4']['IPV'], 2)}."])
    doc.caixa("CUIDADO", "É associação, não efeito causal. O IEG é, porém, o indicador comportamental mais simples de acompanhar (entrega de tarefas).")

    # Q4
    doc.titulo("Q4 — Autoavaliação (IAA)")
    doc.pergunta("As percepções dos alunos sobre si mesmos (IAA) são coerentes com seu desempenho real (IDA) e engajamento (IEG)?")
    al = q4["alinhamento"]
    doc.resposta("A associação com IDA e IEG é positiva, mas fraca. O IAA traz informação diferente da nota e do engajamento.")
    doc.par("O IAA é calculado a partir de seis perguntas sobre como o aluno se sente consigo mesmo, com os estudos, a família, os amigos, "
            "a Associação e os professores (PEDE, Tabela 40). Não é uma previsão da própria nota. Por isso, a diferença IAA − IDA compara "
            "duas escalas distintas e não permite concluir que o aluno “superestima” sua competência.")
    doc.figura("q4_iaa_menos_ida", "Diferença entre IAA e IDA (alunos com IAA > 0; comparação exploratória, corte de 2 pontos).")
    doc.tabela(["Ano", "n", "r (IAA, IDA)", "r (IAA, IEG)", "IAA > IDA + 1", "IAA > IDA + 2", "IAA > IDA + 3"],
               [[x, al[x]["n"], br(q4["corr"][x]["IAA_IDA"]["r"], 3), br(q4["corr"][x]["IAA_IEG"]["r"], 3)]
                + [pct(q4["sensibilidade_corte"][x][c]) for c in ["1", "2", "3"]] for x in anos],
               larguras=[1.6, 1.6, 2.4, 2.4, 2.8, 2.8, 2.8],
               legenda="Associação e diferença entre IAA e IDA, com sensibilidade ao corte (IAA > 0)")
    doc.caixa("LEITURA", f"Em todos os anos o IAA fica, em mediana, cerca de {br(min(al[x]['gap_mediano'] for x in anos), 1)} a "
              f"{br(max(al[x]['gap_mediano'] for x in anos), 1)} pontos acima do IDA, e a parcela com diferença maior que o corte varia muito "
              "conforme o corte escolhido. Para avaliar coerência acadêmica de fato, seria necessário usar só as perguntas sobre estudos.")

    # Q5
    doc.titulo("Q5 — Aspectos psicossociais (IPS)")
    doc.pergunta("Há padrões psicossociais (IPS) que antecedem quedas de desempenho acadêmico ou de engajamento?")
    ant = q5["antecedencia"]
    doc.resposta("Não foi detectada associação entre o IPS de um ano e quedas de IDA ou IEG no ano seguinte, no recorte e método usados.")
    doc.figura("q5_ips_antecedencia", "Percentual de alunos com queda ≥ 1 ponto no ano seguinte, por grupo de IPS no ano anterior.")
    doc.tabela(["Transição", "n", "Corte IPS baixo", "Queda IDA: IPS baixo | demais", "RC queda IDA (IC 95%)", "RC queda IEG (IC 95%)"],
               [[k.replace("->", "→"), v["n"], f"≤ {br(v['corte_ips_baixo'], 2)}",
                 f"{pct(v['taxas_por_grupo']['IPS baixo']['queda_IDA'])} | {pct(v['taxas_por_grupo']['Demais']['queda_IDA'])}",
                 f"{br(v['or_queda_IDA_por_dp_IPS']['or'], 2)} ({ic(v['or_queda_IDA_por_dp_IPS']['ic95'], 2)})",
                 f"{br(v['or_queda_IEG_por_dp_IPS']['or'], 2)} ({ic(v['or_queda_IEG_por_dp_IPS']['ic95'], 2)})"] for k, v in ant.items()],
               larguras=[2.2, 1.4, 2.4, 3.6, 3.4, 3.4],
               legenda="IPS e quedas no ano seguinte. RC = razão de chances por 1 desvio-padrão de IPS no ano (1 = sem associação)")
    doc.caixa("CUIDADO", "Não detectar associação não prova que o aspecto emocional seja irrelevante. O IPS tem poucos valores distintos e "
              "distribuição muito diferente em 2023, o que limita a análise. Os intervalos de confiança incluem 1 nos dois períodos.")

    # Q6
    doc.titulo("Q6 — Aspectos psicopedagógicos (IPP)")
    doc.pergunta("As avaliações psicopedagógicas (IPP) confirmam ou contradizem a defasagem identificada pelo IAN?")
    e23, e24 = q6["em_fase_vs_defasados"]["2023"], q6["em_fase_vs_defasados"]["2024"]
    pb = q6["pct_defasado_por_ipp_baixo"]["2024"]
    doc.resposta("Apontam na mesma direção, com diferença pequena: o IPP é ligeiramente menor entre alunos defasados.")
    doc.figura("q6_ipp_por_ian", "IPP médio por classificação de defasagem (IPP não existe em 2022).")
    doc.itens([
        f"2023: IPP {br(e23['ipp_em_fase'], 2)} em fase (n = {e23['n_em_fase']}) vs {br(e23['ipp_defasados'], 2)} defasados (n = {e23['n_defasados']}), "
        f"Mann-Whitney {fmt_p(e23['p_mannwhitney'])}.",
        f"2024: {br(e24['ipp_em_fase'], 2)} (n = {e24['n_em_fase']}) vs {br(e24['ipp_defasados'], 2)} (n = {e24['n_defasados']}), {fmt_p(e24['p_mannwhitney'])}.",
        f"Correlação IPP × IAN: {br(q6['corr']['2023']['pearson'], 3)} (2023) e {br(q6['corr']['2024']['pearson'], 3)} (2024).",
        f"Em 2024, {pct(pb['ipp_menor_6'])} dos alunos com IPP < 6 (n = {pb['n_ipp_menor_6']}) estão defasados, contra "
        f"{pct(pb['ipp_6_ou_mais'])} dos demais (n = {pb['n_ipp_6_ou_mais']}).",
        f"A classe Severa tem poucos alunos (2024: n = {int(q6['ipp_por_classe']['2024']['Severa']['count'])}); o teste de Kruskal-Wallis "
        f"de 2024 inclui apenas {', '.join(q6['kruskal']['2024']['classes_incluidas'])} (classes com n ≥ 5).",
    ])
    doc.caixa("LEITURA", "O IAN é administrativo (fase × idade) e o IPP é o olhar pedagógico. Eles não se contradizem, mas a diferença é "
              "pequena demais para o IPP identificar a defasagem sozinho.")

    # Q7
    doc.titulo("Q7 — Ponto de virada (IPV)")
    doc.pergunta("Quais comportamentos — acadêmicos, emocionais ou de engajamento — mais influenciam o IPV ao longo do tempo?")
    doc.resposta("Os indicadores mais associados ao IPV são os acadêmicos e de engajamento (IDA e IEG) e, a partir de 2023, o IPP. "
                 "IPS e IAA têm associação pequena.")
    doc.figura("q7_ipv_associacoes", "Coeficientes padronizados de regressão linear do IPV sobre os demais indicadores, por ano.")
    doc.tabela(["Indicador"] + anos, [[ind] + [br(q7["coef_padronizados"][x][ind], 2) if ind in q7["coef_padronizados"][x] else "—" for x in anos]
                                       for ind in ["IDA", "IEG", "IPP", "IAA", "IPS", "IAN"]] +
               [["R²"] + [br(q7["r2"][x], 2) for x in anos], ["n"] + [q7["n"][x] for x in anos]], larguras=[3, 3, 3, 3],
               legenda="Coeficientes padronizados (associação; IAA sem zeros; IPP não medido em 2022)")
    s = q7["corr_ano_seguinte"]["2023->2024"]
    doc.itens([f"Ao longo do tempo, o IPV de 2024 se correlaciona mais com IPV ({br(s['IPV'], 2)}), IDA ({br(s['IDA'], 2)}), "
               f"IPP ({br(s['IPP'], 2)}) e IEG ({br(s['IEG'], 2)}) de 2023 do que com IPS ({br(s['IPS'], 2)}) ou IAA ({br(s['IAA'], 2)}).",
               "Os modelos diferem entre anos (sem IPP em 2022), então os coeficientes não são diretamente comparáveis de um ano para o outro."])
    doc.caixa("CUIDADO", "Os coeficientes descrevem associação, não influência causal. IPV e IPP vêm da avaliação da mesma equipe de "
              "educadores, o que ajuda a explicar a relação forte entre eles.")

    # Q8
    doc.titulo("Q8 — Multidimensionalidade (INDE)")
    doc.pergunta("Quais combinações de indicadores (IDA + IEG + IPS + IPP) elevam mais a nota global do aluno (INDE)?")
    vf = q8["verificacao_formula"]["2024"]
    pa = q8["perfis"]["2024"]["pares_altos"]
    melhor_par = max(pa, key=lambda k: pa[k]["inde"])
    doc.resposta(f"Como o INDE é uma soma ponderada, cada combinação soma os pesos. Entre as quatro dimensões, IDA e IEG pesam 20% cada; "
                 f"IPS e IPP, 10%. O par {melhor_par.replace('+', ' + ')} na mediana ou acima tem o maior INDE médio.")
    doc.par(f"A fórmula oficial (INDE = 0,1·IAN + 0,2·IDA + 0,2·IEG + 0,1·IAA + 0,1·IPS + 0,1·IPP + 0,2·IPV) reproduz o INDE de 2024 "
            f"com R² = {br(vf['r2'], 3)} (n = {vf['n']}).")
    doc.figura("q8_potencial_aritmetico_inde", "Potencial aritmético: ganho de INDE se cada indicador chegasse a 10 (média de 2024, fases 0–7).")
    doc.figura("q8_inde_por_numero_de_dimensoes", "INDE médio pelo número de dimensões (IDA, IEG, IPS, IPP) na mediana do ano ou acima.")
    doc.itens([
        "Cada ponto a mais em IDA, IEG ou IPV soma 0,2 ao INDE; em IPS ou IPP, 0,1.",
        f"Alunos com as quatro dimensões na mediana ou acima têm INDE médio {br(perfil['4']['mean'], 2)} (n = {int(perfil['4']['count'])}); "
        f"com nenhuma, {br(perfil['0']['mean'], 2)} (n = {int(perfil['0']['count'])}) (2024).",
        f"Pares (2024): {melhor_par.replace('+', ' + ')} = {br(pa[melhor_par]['inde'], 2)} (n = {pa[melhor_par]['n']}); "
        f"IPS + IPP = {br(pa['IPS+IPP']['inde'], 2)} (n = {pa['IPS+IPP']['n']}).",
    ])
    doc.caixa("LEITURA", f"O {maior_potencial} tem o maior espaço aritmético de ganho porque combina peso alto e média baixa. Isso indica onde "
              "há mais margem no índice, não qual intervenção traria mais retorno: essa comparação exigiria custos, viabilidade e efeito estimado das ações.")

    # Q9
    doc.quebra()
    doc.titulo("Q9 — Previsão de risco com Machine Learning")
    doc.pergunta("Quais padrões nos indicadores permitem identificar alunos em risco antes de queda no desempenho ou aumento da defasagem? "
                 "Construa um modelo preditivo que mostre a probabilidade de o aluno entrar em risco de defasagem.")
    doc.resposta(f"O modelo estima o risco de redução de D (fase efetiva − fase ideal) no ano seguinte, com ROC-AUC {br(t['roc_auc'], 3)} "
                 f"(IC 95% {ic(t['roc_auc_ic95'])}) no teste temporal. Os sinais mais associados são idade, defasagem atual e IPV.")
    doc.titulo("Evento previsto e população", "Heading 2")
    eventos = {k: sum(v.get(k, 0) for v in dec.values()) for k in ["Entrada em defasagem", "Agravamento (já defasado)", "Perda de vantagem (continua em fase)"]}
    doc.par(f"Evento: D diminui de um ano para o outro, ou seja, a adequação de nível piora (aumento da defasagem na definição da PEDE). "
            f"População: alunos das fases 0 a 7 presentes em dois anos consecutivos; a probabilidade é condicional a o aluno aparecer no ano seguinte. "
            f"Prevalência: {pct(m['prevalencia_treino'] * 100)} (2022→23) e {pct(m['prevalencia_teste'] * 100)} (2023→24).")
    doc.itens([
        ("O evento reúne três situações: ", f"entrada em defasagem ({eventos['Entrada em defasagem']} casos), agravamento de quem já estava "
         f"defasado ({eventos['Agravamento (já defasado)']}) e perda de vantagem de quem continua em fase ({eventos['Perda de vantagem (continua em fase)']})."),
        ("Relação com a promoção: ", f"quase todo evento ocorre entre não promovidos ({pxe['promovidos_com_evento']} entre {pxe['promovidos']} promovidos), "
         f"mas {pxe['nao_promovidos_sem_evento']} dos {pxe['nao_promovidos']} não promovidos ({pct(pxe['nao_promovidos_sem_evento'] / pxe['nao_promovidos'] * 100)}) "
         "não tiveram o evento. O evento não equivale à retenção."),
    ])
    doc.titulo("Método", "Heading 2")
    doc.itens([
        ("Separação temporal: ", f"treino com atributos de 2022 e evento em 2023 ({m['pares_treino']} alunos); teste com 2023 → 2024 ({m['pares_teste']})."),
        ("Atributos (fechamento do ano t): ", "fase, idade, idade relativa à fase, defasagem, anos no programa, gênero, tipo de escola, IDA, IEG, IAA, "
         "IPV, Matemática, Português e INDE. IPP (ausente em 2022) e IPS (distribuição de 2023 atípica) ficaram de fora."),
        ("Seleção e calibração: ", f"cinco candidatos comparados por validação cruzada só no treino; escolhido: {m['modelo_escolhido']}. "
         "A calibração sigmoide foi definida antes de consultar o teste, porque o enunciado pede uma probabilidade."),
        ("Limiares: ", f"definidos no treino. Alto ≥ {br(m['limiar_alto'], 3)} (máximo F1); médio ≥ {br(m['limiar_medio'], 3)} (captura 85% dos eventos)."),
        ("Transparência: ", "a formulação do evento foi escolhida após explorar alvos alternativos no período 2023→24. O teste é independente "
         "para modelo, calibração e limiares, mas não totalmente para a formulação; por isso também se reporta a direção inversa."),
    ])
    doc.titulo("Resultados no teste (2023 → 2024)", "Heading 2")
    linhas = [[r["modelo"], br(r["ROC-AUC"], 3), br(r["AP"], 3), "—" if r["Brier"] is None else br(r["Brier"], 3)] for r in m["teste"]]
    doc.tabela(["Modelo", "ROC-AUC", "AP", "Brier"], linhas, larguras=[8, 2.6, 2.6, 2.6],
               legenda=f"Desempenho no teste (n = {m['pares_teste']}). AP = average precision; referência sem informação = prevalência")
    doc.par(f"Intervalos de confiança (bootstrap, 2.000 reamostragens): ROC-AUC {ic(t['roc_auc_ic95'])}; AP {ic(t['ap_ic95'])}. "
            f"Na direção inversa (treino 2023→24, teste 2022→23), ROC-AUC {br(rob['Direção inversa (treina 23→24, testa 22→23)']['ROC-AUC'], 3)}. "
            "Os intervalos são condicionais ao ano de teste e não medem estabilidade em anos futuros.")
    doc.figura("q9_roc_calibracao", "Curva ROC e calibração no teste (antes e depois da calibração ajustada no treino).")
    doc.tabela(["Faixa", "Alunos sinalizados", "Precisão", "Recall"], [
        ["Alto", ma["alunos sinalizados"], pct(ma["precisão"] * 100), pct(ma["recall"] * 100)],
        ["Médio ou alto", mm["alunos sinalizados"], pct(mm["precisão"] * 100), pct(mm["recall"] * 100)],
    ], larguras=[4, 4, 3, 3], legenda=f"Uso operacional no teste ({m['pares_teste']} alunos)")
    doc.itens([
        f"Acompanhando os {top10['acompanhar top']} com maior risco, a equipe encontra {top10['eventos capturados']} eventos "
        f"({pct(top10['% do total de eventos'])} do total), {br(top10['lift vs aleatório'], 1)} vezes a taxa de uma escolha aleatória; "
        f"com os {top20['acompanhar top']}, {pct(top20['% do total de eventos'])}.",
        f"Composição da faixa Alto no teste (n = {n_alto}): {comp_alto.get('Entrada em defasagem', 0)} entradas em defasagem, "
        f"{comp_alto.get('Agravamento (já defasado)', 0)} agravamentos, {comp_alto.get('Perda de vantagem (continua em fase)', 0)} perdas de vantagem "
        f"e {comp_alto.get('Sem evento', 0)} sem evento.",
        "Calibração por fase: a maior diferença entre previsto e observado ocorre em "
        + " e ".join(f"{r['fase']} ({br(r['previsto'] * 100, 0)}% previsto vs {br(r['observado'] * 100, 0)}% observado, n = {r['alunos']})" for r in pior_calib)
        + ", onde a promoção caiu entre os anos.",
    ])
    doc.figura("q9_importancia", "Importância por permutação no teste (queda de ROC-AUC ao embaralhar cada atributo).")
    doc.titulo("Sensibilidades", "Heading 2")
    doc.tabela(["Análise", "ROC-AUC (teste)"],
               [[k, br(v["ROC-AUC (teste)"], 3)] for k, v in sens.items()]
               + [["Rótulos de 2024 pela regra idade → fase de 2022/23", f"{br(s_regra['roc_auc_rotulos_regra'], 3)} (IC {ic(s_regra['ic95_auc_rotulos_regra'])})"],
                  [f"Só alunos novos (n = {rob['Só alunos novos (fora do treino)']['n']})", br(rob["Só alunos novos (fora do treino)"]["ROC-AUC"], 3)],
                  [f"Só alunos já vistos no treino (n = {rob['Só alunos já vistos no treino']['n']})", br(rob["Só alunos já vistos no treino"]["ROC-AUC"], 3)]],
               larguras=[11, 4], legenda="Sensibilidade do modelo a decisões de dados e de formulação (nenhuma usada para escolher o modelo)")
    doc.itens([
        f"As decisões sobre IAA, INDE e IPS mudam pouco a ordenação (ROC-AUC entre {br(min(faixa_teste), 3)} e {br(max(faixa_teste), 3)}).",
        f"Com a regra de fase ideal de 2022/23, {s_regra['rotulos_alterados_teste']} rótulos do teste mudam "
        f"({s_regra['eventos_fonte']} → {s_regra['eventos_regra']} eventos) e a ROC-AUC cai para {br(s_regra['roc_auc_rotulos_regra'], 3)}. "
        "É um cenário alternativo, não uma correção.",
        f"O alvo da versão anterior, com os mesmos atributos e protocolo, atinge ROC-AUC {br(anterior['ROC-AUC (teste)'], 3)}.",
    ])
    doc.titulo("Modelo final", "Heading 2")
    fx = final["alunos_2024_por_faixa"]
    doc.par(f"O modelo usado na aplicação é reajustado com as duas transições ({final['cv_agrupada']['n']} pares), com calibração e limiares "
            f"obtidos por validação cruzada agrupada por aluno (ROC-AUC {br(final['cv_agrupada']['roc_auc'], 3)}). Ele não tem teste futuro. "
            f"Aplicado aos {final['n_alunos_2024']} alunos de 2024 das fases 0–7, classifica {fx.get('Alto', 0)} como Alto, {fx.get('Médio', 0)} como Médio "
            f"e {fx.get('Baixo', 0)} como Baixo. A proporção em Alto é maior que no teste, em parte porque a regra de 2024 deixa mais alunos em fase, "
            "o que os coloca em risco de perder essa condição.")
    doc.caixa("USO RECOMENDADO", "Usar o ranking e as faixas para priorizar acompanhamento, com os indicadores do fechamento do ano anterior. "
              "Não ler a probabilidade como diagnóstico individual, principalmente nas Fases 5 e 7. O modelo não estima risco de saída da base "
              "nem prova causas: a importância dos atributos é associação preditiva.")

    # Q10
    doc.quebra()
    doc.titulo("Q10 — Efetividade do programa (Pedras)")
    doc.pergunta("Os indicadores mostram melhora consistente ao longo do ciclo nas diferentes fases (Quartzo, Ágata, Ametista e Topázio), confirmando o impacto real do programa?")
    doc.resposta("Não de forma consistente. A distribuição das Pedras oficiais melhora, mas, acompanhando os mesmos alunos, o INDE fica "
                 "estável, exceto no grupo que começou como Quartzo. Sem grupo de comparação, os dados não permitem medir impacto causal.")
    doc.figura("q10_distribuicao_pedras", "Distribuição das Pedras oficiais por ano (alunos com Pedra atribuída).")
    lim = q10["limites_observados"]
    doc.tabela(["Ano"] + ["Ágata (máx.)", "Ametista (mín.–máx.)", "Topázio (mín.)", "Divergências com cortes 6/7/8"],
               [[x, br(lim[f"{x}|Ágata"]["max"], 3), f"{br(lim[f'{x}|Ametista']['min'], 3)}–{br(lim[f'{x}|Ametista']['max'], 3)}",
                 br(lim[f"{x}|Topázio"]["min"], 3), q10["divergencias_regra_6_7_8"][x]] for x in anos],
               larguras=[2, 3, 4, 3, 4.5], legenda="Limites de INDE observados em cada Pedra oficial (não são os cortes oficiais)")
    doc.figura("q10_coorte_inde_por_pedra", f"INDE médio dos mesmos {q10['coorte_n']} alunos nos três anos, agrupados pela Pedra de 2022.")
    doc.tabela(["Transição", "Alunos", "Avançou", "Permaneceu", "Recuou", "Variação média do INDE"],
               [[k.replace("->", "→"), v["n"], pct(v["avancou"]), pct(v["permaneceu"]), pct(v["recuou"]), br(v["delta_inde_medio"], 2)]
                for k, v in q10["transicoes"].items()],
               larguras=[2.8, 2, 2.4, 2.6, 2.4, 3.6], legenda="Mobilidade entre Pedras oficiais de cada ano (mesmos alunos)")
    doc.itens([
        f"Topázio passa de {pct(q10['distribuicao_pct']['2022']['Topázio'])} para {pct(q10['distribuicao_pct']['2024']['Topázio'])} dos alunos; "
        f"Quartzo, de {pct(q10['distribuicao_pct']['2022']['Quartzo'])} para {pct(q10['distribuicao_pct']['2024']['Quartzo'])}.",
        f"Os limites das Pedras em 2022 não coincidem com os de 2023 e 2024: aplicar cortes 6/7/8 a 2022 mudaria "
        f"{q10['divergencias_regra_6_7_8']['2022']} classificações. A mobilidade entre Pedras mistura mudança de INDE e possível mudança de critério.",
        "Nas transições anuais, avanços e recuos se equilibram, e a variação média do INDE dos mesmos alunos fica próxima de zero.",
        f"Na coorte, só o grupo Quartzo sobe ({br(coorte['Quartzo']['2022'], 2)} → {br(coorte['Quartzo']['2024'], 2)}, n = {coorte['Quartzo']['n']}); "
        "parte dessa subida, e da oscilação do Topázio, é esperada por regressão à média.",
        "Os alunos que deixam de aparecer na base são, em média, de Pedras mais baixas (Q11), então a distribuição agregada também reflete mudança de composição.",
    ])
    doc.caixa("LEITURA", "O resultado observado é estabilidade do INDE entre os alunos que permanecem, com melhora descritiva no grupo de "
              "partida mais baixo. Atribuir isso ao programa exigiria um grupo de comparação.")

    # Q11
    doc.titulo("Q11 — Insights adicionais e recomendações")
    ev = nr["2022->2023"]
    doc.titulo("1. Não reaparecimento na base", "Heading 2")
    doc.figura("q11_nao_reaparecimento_por_pedra", "Percentual de alunos ausentes na base do ano seguinte, por Pedra.")
    doc.par(f"{pct(ev['pct'])} dos alunos de 2022 não aparecem na base de 2023; o INDE médio desses alunos era {br(ev['inde_ausentes'], 2)}, "
            f"contra {br(ev['inde_presentes'], 2)} dos que aparecem. A base não diz se a ausência é desligamento, conclusão, interrupção ou falta de "
            f"avaliação, e {nr['ausentes_2023_que_reaparecem_2024']} dos {nr['ausentes_2023']} ausentes reaparecem em 2024. "
            "Por isso, isto não é uma taxa de evasão comprovada. É um sinal a investigar com os registros de matrícula.")
    doc.tabela(["Transição"] + list(nrp["2022->2023"]["pct"].keys()),
               [[k.replace("->", "→")] + [f"{pct(v['pct'][p])} (n = {v['n'][p]})" for p in v["pct"]] for k, v in nrp.items()],
               larguras=[2.8, 3.4, 3.4, 3.4, 3.4], legenda="Não reaparecimento na base do ano seguinte, por Pedra")
    doc.titulo("2. Promoção de fase", "Heading 2")
    doc.figura("q11_taxa_promocao", "Taxa de promoção à fase seguinte entre alunos presentes nos dois anos (fases 0–7).")
    doc.par(f"A promoção é mais baixa no ALFA ({pct(tp['0']['2022'], 0)} e {pct(tp['0']['2023'], 0)}). Em 2023→24 caiu na Fase 5 "
            f"({pct(tp['5']['2022'], 0)} → {pct(tp['5']['2023'], 0)}; n = {npromo['5']['2023']}) e na Fase 7 "
            f"({pct(tp['7']['2022'], 0)} → {pct(tp['7']['2023'], 0)}; n = {npromo['7']['2023']}). "
            "Como quase todo evento de redução de D ocorre entre não promovidos, essas fases concentram o risco.")
    doc.titulo("3. O INDE sozinho não ordena o risco", "Heading 2")
    pq = q11["evento_por_quartil_inde"]
    doc.par("O percentual de alunos com redução de D é parecido nos quatro quartis de INDE ("
            + ", ".join(f"Q{k}: {pct(v['pct'])}, n = {v['n']}" for k, v in pq.items())
            + f"), e o INDE, usado sozinho, tem ROC-AUC {br(q11['auc_inde_sozinho'], 3)} para esse evento. Isso não prova que o risco independa "
            "do INDE quando se consideram outros fatores, mas mostra que acompanhar apenas a Pedra não identifica esses alunos.")
    doc.titulo("Recomendações à Passos Mágicos", "Heading 2")
    doc.itens([
        ("Priorização preventiva: ", "com os indicadores do fechamento de cada ano, aplicar o modelo e acompanhar primeiro os alunos da faixa Alto, "
         "com atenção ao ALFA e ao Ensino Médio."),
        (f"Reforço acadêmico na {fase(min(menor_sempre))}: ", "é a fase com menor IDA nos três anos."),
        ("Investigar ausências na base: ", "cruzar com matrícula, conclusão e desligamento para saber quanto é evasão e por quê."),
        ("Documentar critérios: ", "registrar a regra de fase ideal, os cortes de Pedra e os critérios de promoção de cada ano."),
        ("Qualidade dos dados: ", "padronizar códigos entre anos, distinguir zero de não avaliado (IEG, IAA), manter a escala do IPS e registrar o IPP em todas as fases."),
        ("Autoavaliação: ", "separar as perguntas sobre estudos das demais, para comparar percepção e desempenho acadêmico de forma válida."),
    ])

    doc.titulo("Cuidados metodológicos")
    doc.itens([
        "Os resultados são observacionais: associação não significa causalidade.",
        "A base tem uma fotografia por ano; a composição muda entre anos (entradas e ausências). Por isso foram incluídas análises dos mesmos alunos.",
        "A regra de fase ideal de 2024 difere da de 2022/23; resultados que dependem da defasagem trazem a sensibilidade correspondente.",
        "IAN, Pedra e INDE são derivados de outros indicadores; relações entre eles são parcialmente estruturais.",
        "Valores ausentes não foram tratados como zero; os tratamentos analíticos (IAA, IEG) são hipóteses documentadas, com a fonte preservada.",
    ])
    p = doc.doc.add_paragraph()
    r = p.add_run("Documento gerado por src/documento.py a partir dos resultados do pipeline (outputs/resultados_analises.json e outputs/resultados_modelo.json).")
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = CINZA
    return doc


def carregar_json(caminho):
    with open(caminho, encoding="utf-8") as f:
        return json.load(f, parse_constant=lambda c: (_ for _ in ()).throw(ValueError(f"JSON não estrito: {c}")))


def main():
    construir(carregar_json(RESULTADOS_ANALISES), carregar_json(RESULTADOS_MODELO)).salvar(SAIDA)
    print(f"Documento: {SAIDA}")


if __name__ == "__main__":
    main()
